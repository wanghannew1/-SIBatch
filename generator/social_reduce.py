"""generator/social_reduce.py — T9 社保减员报盘生成器（.xlsx 5 列 + 追加合同终止原因）。

输出「单位职工停保模板」：Sheet1 表头 1 列名（1-based 行 1），数据自 1-based 行 2。
列：证件号码 / 姓名 / 停保日期 / 停保原因 / 备注，末尾追加第 6 列「合同终止原因」。

数据来源：reduce_records 中 social_reduce_ym 非空的记录（仅社保减员）。
停保原因不硬编码 — 读 DB 表 mapping_stop_reason（source_reason = 合同终止原因）。
"""

from __future__ import annotations

import os

from columns import (
    SOCIAL_REDUCE_OUTPUT_SHEET,
    SOCIAL_REDUCE_OUTPUT_HEADER_ROW_1BASED,
    SOCIAL_REDUCE_OUTPUT_DATA_START_ROW_1BASED,
    SOCIAL_REDUCE_OUTPUT_HEADERS,
    SOCIAL_REDUCE_OUTPUT_APPEND_COLUMN,
)
from db import query_all, query_one
from mapping import dash_date_to_yyyymmdd

HEADERS = tuple(SOCIAL_REDUCE_OUTPUT_HEADERS) + (SOCIAL_REDUCE_OUTPUT_APPEND_COLUMN,)

_SELECT_RECORDS = (
    "SELECT * FROM reduce_records WHERE batch_id = ? AND social_reduce_ym IS NOT NULL"
    " AND TRIM(social_reduce_ym) <> '' ORDER BY id"
)

_SELECT_STOP_REASON = (
    "SELECT social_stop_reason FROM mapping_stop_reason WHERE source_reason = ?"
)


def _stop_reason(terminate_reason):
    """合同终止原因 → 社保停保原因（读 DB 映射）；无映射时回退原始值。"""
    row = query_one(_SELECT_STOP_REASON, (terminate_reason,))
    if row is not None and row["social_stop_reason"]:
        return row["social_stop_reason"]
    return terminate_reason or ""


def generate(batch_id, out, *, generated_by=None):
    """生成社保减员报盘 .xlsx，返回输出文件路径。

    batch_id      — import_batches.id（减员批次）
    out           — 输出 .xlsx 路径（新建文件）
    generated_by  — 生成人（保留参数，供调用方记录）
    """
    rows = query_all(_SELECT_RECORDS, (batch_id,))

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = SOCIAL_REDUCE_OUTPUT_SHEET

    header_row = SOCIAL_REDUCE_OUTPUT_HEADER_ROW_1BASED
    data_start = SOCIAL_REDUCE_OUTPUT_DATA_START_ROW_1BASED
    for col_idx, name in enumerate(HEADERS, start=1):
        ws.cell(row=header_row, column=col_idx, value=name)

    for offset, row in enumerate(rows):
        r = data_start + offset
        ws.cell(row=r, column=1, value=row["id_card"])
        ws.cell(row=r, column=2, value=row["name"])
        ws.cell(row=r, column=3, value=dash_date_to_yyyymmdd(row["terminate_date"]) or "")
        ws.cell(row=r, column=4, value=_stop_reason(row["terminate_reason"]))
        ws.cell(row=r, column=5, value=row["remark"] or "")
        ws.cell(row=r, column=6, value=row["terminate_reason"] or "")

    out_dir = os.path.dirname(os.path.abspath(out))
    os.makedirs(out_dir, exist_ok=True)
    wb.save(out)
    return out
