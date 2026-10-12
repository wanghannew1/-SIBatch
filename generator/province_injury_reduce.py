"""省工伤减员报盘生成器（T10）。

数据来源：`province_injury` 表（独立维护的省工伤人员信息表），
筛选需减员的记录（停保日期非空）。

输出：新建 `.xlsx`，列名与「单位职工停保模板」一致（5 列）：
    证件号码 / 姓名 / 停保日期 / 停保原因 / 备注
    —— 不追加「合同终止原因」列（省工伤名单无该字段）。

列映射：
    证件号码 = id_card
    姓名     = name
    停保日期 = stop_date（含 '-' 时归一化为 yyyymmdd）
    停保原因 = stop_reason（原始值，不做映射）
    备注     = ''
"""

import os

from openpyxl import Workbook

import db
from columns import SOCIAL_REDUCE_OUTPUT_HEADERS
from mapping import dash_date_to_yyyymmdd

HEADERS = tuple(SOCIAL_REDUCE_OUTPUT_HEADERS)

_SOURCE_SQL = (
    "SELECT * FROM province_injury WHERE id IN ({placeholders})"
    " AND stop_date IS NOT NULL AND TRIM(stop_date) <> '' ORDER BY id"
)


def _normalize_stop_date(value):
    """停保日期归一化：'2026-07-01' → '20260701'；已紧凑则原样返回。"""
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    normalized = dash_date_to_yyyymmdd(text)
    return normalized if normalized is not None else text


def generate(record_ids, out, *, generated_by=None):
    """按选中人员生成省工伤减员报盘 .xlsx，返回输出文件路径。

    record_ids 为 province_injury.id 列表（仅停保人员可生成）。
    generated_by 为操作人标识，预留用于生成审计（output_batches）。
    """
    if not record_ids:
        return out
    sql = _SOURCE_SQL.format(placeholders=",".join("?" * len(record_ids)))
    rows = db.query_all(sql, tuple(record_ids))

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(list(HEADERS))
    for row in rows:
        ws.append(
            [
                row["id_card"],
                row["name"],
                _normalize_stop_date(row["stop_date"]),
                row["stop_reason"] or "",
                "",
            ]
        )

    parent = os.path.dirname(os.path.abspath(out))
    if parent:
        os.makedirs(parent, exist_ok=True)
    wb.save(out)
    return out
