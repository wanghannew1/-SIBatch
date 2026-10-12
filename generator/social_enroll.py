"""T8 — 社保增员报盘生成器（企业职工批量参保登记模板，人员名单 27 列 A:AA）。

以真实模板为底稿（openpyxl 打开 → 从 1-based 第 5 行写数据 → 另存），
完整保留 R1 说明 / R2 单位信息 / R3 列名 / R4「小明」示例行 / 合并单元格。
所有字典与默认值均读库（dict_mapping / system_defaults / unit_config），
业务规则函数取自 mapping.py，列结构取自 columns.py。
"""

import json
import os

import openpyxl

import columns
import db
from mapping import dash_date_to_yyyymmdd, derive_birthday, derive_gender, strip_prefix

_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

_UNIT_CODE_CELL = (2, 2)    # B2 — 「单位编号:」标签(A2)后的值单元格
_UNIT_NAME_CELL = (2, 5)    # E2 — 「单位名称:」标签(C2:D2 合并)后的值单元格(E2:L2 合并区)


def _template_path():
    return os.path.join(_BASE_DIR, columns.SOCIAL_ENROLL_OUTPUT_FILE)


def _load_dict():
    rows = db.query_all("SELECT category, source_code, target_value FROM dict_mapping")
    return {(r["category"], r["source_code"]): r["target_value"] for r in rows}


def _load_defaults():
    rows = db.query_all("SELECT key, value FROM system_defaults")
    return {r["key"]: r["value"] for r in rows}


def _lookup(mapping, category, *keys):
    for key in keys:
        if key is None:
            continue
        value = mapping.get((category, key))
        if value:
            return value
    return None


def _yes_no(value):
    text = "" if value is None else str(value)
    return "是" if ("是" in text or text.startswith("1")) else "否"


def _address(raw_json):
    if not raw_json:
        return ""
    try:
        data = json.loads(raw_json)
    except (TypeError, ValueError):
        return ""
    if not isinstance(data, dict):
        return ""
    address = data.get("地址", "")
    return "" if address is None else str(address)


def _unit_config():
    return db.query_one(
        "SELECT unit_code, unit_name FROM unit_config WHERE insurance_type = 'social'"
    )


def _record_values(record, mapping, defaults):
    ethnicity = record["ethnicity"]
    person_identity = record["person_identity"]
    social_location = record["social_location"]
    doc_type = record["doc_type"]
    address = _address(record["raw_json"])
    doc_type_key = strip_prefix(doc_type) if doc_type else ""
    contract_start = dash_date_to_yyyymmdd(record["contract_start"])
    return (
        _lookup(mapping, "doc_type", doc_type_key, doc_type) or "",
        record["id_card"],
        record["name"],
        _lookup(mapping, "ethnicity", ethnicity, strip_prefix(ethnicity))
        or strip_prefix(ethnicity)
        or "",
        derive_gender(record["id_card"]) or "",
        derive_birthday(record["id_card"]) or "",
        contract_start or "",
        defaults.get("国籍") or "",
        defaults.get("农民工标识") or "",
        _lookup(mapping, "person_identity", person_identity, strip_prefix(person_identity))
        or strip_prefix(person_identity)
        or "",
        defaults.get("临时账户标识") or "",
        record["contract_salary"],
        contract_start or "",
        _yes_no(record["participate_pension"]),
        _yes_no(record["participate_unemployment"]),
        _yes_no(record["participate_injury"]),
        defaults.get("是否兼职工伤") or "",
        address,
        address,
        _lookup(mapping, "location_work_address", social_location, strip_prefix(social_location))
        or "",
        record["phone"],
        defaults.get("参保身份") or "",
        None,
        None,
        None,
        None,
        defaults.get("模板版本") or "",
    )


def generate(record_ids, out, *, generated_by=None):
    """按选中人员生成社保增员报盘，返回输出路径。

    ``record_ids`` 为 enroll_records.id 列表（跨批次）。只取 social_location
    非空的记录（市社保报盘；省工伤不走此报盘）。unit_config(social) 为 NULL
    时 R2 保持模板原样，不臆造值。
    """
    mapping = _load_dict()
    defaults = _load_defaults()

    if not record_ids:
        return out

    placeholders = ",".join("?" * len(record_ids))
    records = db.query_all(
        f"SELECT * FROM enroll_records WHERE id IN ({placeholders})"
        " AND social_location IS NOT NULL AND TRIM(social_location) <> ''"
        " ORDER BY id",
        tuple(record_ids),
    )

    wb = openpyxl.load_workbook(_template_path())
    ws = wb[columns.SOCIAL_ENROLL_OUTPUT_SHEET]

    unit = _unit_config()
    if unit is not None:
        if unit["unit_code"] is not None:
            ws.cell(*_UNIT_CODE_CELL).value = unit["unit_code"]
        if unit["unit_name"] is not None:
            ws.cell(*_UNIT_NAME_CELL).value = unit["unit_name"]

    start = columns.SOCIAL_ENROLL_OUTPUT_DATA_START_ROW_1BASED
    for offset, record in enumerate(records):
        values = _record_values(record, mapping, defaults)
        row = start + offset
        for col_idx, value in enumerate(values):
            if value is not None:
                ws.cell(row, col_idx + 1).value = value

    out_dir = os.path.dirname(os.path.abspath(out))
    os.makedirs(out_dir, exist_ok=True)
    wb.save(out)

    return out
